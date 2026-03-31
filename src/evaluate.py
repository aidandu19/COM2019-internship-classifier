# evaluate.py - evaluation pipeline for COM2019 internship classifier
#
# Reads hand-labelled gold standard (153 entries) from xlsx, extracts matching
# LLM + baseline predictions from classified_listings.xlsx (501 entries), and
# compares random baseline vs rule-based baseline vs LLM classifier using:
# - Per-class precision, recall, F1 (sklearn classification_report)
# - Confusion matrix heatmaps
# - Confidence calibration analysis (binned accuracy vs stated confidence)
# - Expected Calibration Error (ECE)
# - Reliability diagrams
# - Error analysis (top confusion pairs per dimension)
# - Class distribution report
# - Unseen split descriptive statistics (348 non-gold entries)

import os
import random
from collections import Counter

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from sklearn.metrics import classification_report, confusion_matrix

from .taxonomy import (
    FIRM_TYPES, ROLE_FUNCTIONS, PROGRAMME_STATUSES,
    FIRM_TYPE_CODES, ROLE_FUNCTION_CODES, PROGRAMME_STATUS_CODES,
)
from .baseline import rule_based_baseline, random_baseline


FIGURES_DIR = os.path.join("output", "figures")
GOLD_STANDARD_PATH = os.path.join("data", "GOLD STANDARD HAND FILLED LISTINGS.xlsx")
CLASSIFIED_PATH = os.path.join("data", "classified_listings.xlsx")

# Calibration bin edges: four bins covering low/medium/high/very-high confidence
CALIBRATION_BINS = [0.0, 0.3, 0.6, 0.8, 1.0]


# ---------------------------------------------------------------------------
# 1. Load gold standard from xlsx
# ---------------------------------------------------------------------------

def _convert_code(code, mapping, valid_list):
    """Convert a short code to its full enum name. Returns None on failure."""
    if code is None or str(code).strip() == "":
        return None
    code_str = str(code).strip().upper()
    if code_str in valid_list:
        return code_str
    return mapping.get(code_str)


def load_gold_standard(path=GOLD_STANDARD_PATH):
    """Load hand-labelled gold standard from xlsx.

    Reads the 'Gold Standard' sheet with columns:
      Col 2: Company, Col 3: Programme, Col 4: Opening, Col 5: Closing,
      Col 6: YOUR_st, Col 7: YOUR_ft, Col 8: YOUR_rf, Col 9: rationale

    Returns list of dicts with full enum names, or None if file missing.
    """
    if not os.path.exists(path):
        print(f"[evaluate] Gold standard not found at {path}")
        return None

    from openpyxl import load_workbook
    wb = load_workbook(path)

    if "Gold Standard" not in wb.sheetnames:
        print(f"[evaluate] Sheet 'Gold Standard' not found. Available: {wb.sheetnames}")
        wb.close()
        return None

    ws = wb["Gold Standard"]
    entries = []
    warnings = []

    for row_idx in range(2, ws.max_row + 1):
        company = ws.cell(row=row_idx, column=2).value
        if company is None or str(company).strip() == "":
            continue

        company = str(company).strip()
        programme = str(ws.cell(row=row_idx, column=3).value or "").strip()
        opening = str(ws.cell(row=row_idx, column=4).value or "").strip()
        closing = str(ws.cell(row=row_idx, column=5).value or "").strip()

        ft_raw = ws.cell(row=row_idx, column=7).value
        st_raw = ws.cell(row=row_idx, column=6).value
        rf_raw = ws.cell(row=row_idx, column=8).value

        ft = _convert_code(ft_raw, FIRM_TYPE_CODES, FIRM_TYPES)
        st = _convert_code(st_raw, PROGRAMME_STATUS_CODES, PROGRAMME_STATUSES)
        rf = _convert_code(rf_raw, ROLE_FUNCTION_CODES, ROLE_FUNCTIONS)

        if ft is None:
            warnings.append(f"  Row {row_idx} ({company}): unrecognised firm_type code '{ft_raw}'")
        if st is None:
            warnings.append(f"  Row {row_idx} ({company}): unrecognised programme_status code '{st_raw}'")
        if rf is None:
            warnings.append(f"  Row {row_idx} ({company}): unrecognised role_function code '{rf_raw}'")

        # Skip entries with any unresolved label
        if ft is None or st is None or rf is None:
            continue

        entries.append({
            "company_name": company,
            "programme_name": programme,
            "opening_date": opening,
            "closing_date": closing,
            "firm_type": ft,
            "programme_status": st,
            "role_function": rf,
        })

    wb.close()

    if warnings:
        print(f"[evaluate] {len(warnings)} label warnings:")
        for w in warnings:
            print(w)

    if len(entries) < 5:
        print(f"[evaluate] Only {len(entries)} valid gold standard entries, need >= 5.")
        return None

    return entries


# ---------------------------------------------------------------------------
# 1b. Load classified results from xlsx (501 entries)
# ---------------------------------------------------------------------------

def load_classified_results(path=CLASSIFIED_PATH):
    """Load pipeline output from classified_listings.xlsx.

    Returns list of dicts with LLM and baseline predictions, or None if missing.

    Column mapping (from src/export.py):
      1: Company Name, 2: Programme, 3-5: dates/stage,
      6: Baseline Firm Type, 7: Baseline Firm Conf, 8: Baseline Role, 9: Baseline Status,
      10: LLM Firm Type, 11: LLM Firm Conf, 12: LLM Firm Rationale,
      13: LLM Role, 14: LLM Role Conf, 15: LLM Role Rationale,
      16: LLM Status, 17: LLM Status Conf, 18: LLM Status Rationale,
      19: LLM Model Used, 20: LLM Escalated
    """
    if not os.path.exists(path):
        print(f"[evaluate] Classified results not found at {path}")
        return None

    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active
    results = []

    for row_idx in range(2, ws.max_row + 1):
        company = ws.cell(row=row_idx, column=1).value
        if company is None or str(company).strip() == "":
            continue

        def _cell(col):
            v = ws.cell(row=row_idx, column=col).value
            return v if v is not None else ""

        def _float(col):
            v = ws.cell(row=row_idx, column=col).value
            try:
                return float(v)
            except (ValueError, TypeError):
                return 0.5

        results.append({
            "company_name": str(company).strip(),
            "programme_name": str(_cell(2)).strip(),
            "opening_date": str(_cell(3)).strip(),
            "closing_date": str(_cell(4)).strip(),
            "latest_stage": str(_cell(5)).strip(),
            # Baseline
            "baseline_firm_type": str(_cell(6)).strip(),
            "baseline_firm_type_confidence": _float(7),
            "baseline_role_function": str(_cell(8)).strip(),
            "baseline_programme_status": str(_cell(9)).strip(),
            # LLM
            "firm_type": str(_cell(10)).strip(),
            "firm_type_confidence": _float(11),
            "firm_type_rationale": str(_cell(12)).strip(),
            "role_function": str(_cell(13)).strip(),
            "role_function_confidence": _float(14),
            "role_function_rationale": str(_cell(15)).strip(),
            "programme_status": str(_cell(16)).strip(),
            "programme_status_confidence": _float(17),
            "programme_status_rationale": str(_cell(18)).strip(),
            "model_used": str(_cell(19)).strip(),
            "escalated": _cell(20),
        })

    wb.close()
    return results


# ---------------------------------------------------------------------------
# 1c. Match gold standard against classified results
# ---------------------------------------------------------------------------

def _normalize_for_matching(text):
    """Normalize text for fuzzy matching across encoding differences.

    The TRACKR CSV uses latin-1 encoding which mangles unicode characters:
    en/em dashes become Ð, accented characters may be lost, etc.
    """
    import unicodedata
    # Normalize unicode (e.g. é -> e)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    # Collapse all dash-like characters to a simple hyphen
    for ch in ("\u2013", "\u2014", "\u00d0", "\u2010", "\u2011", "\u2012"):
        text = text.replace(ch, "-")
    return text.strip().lower()


def match_gold_to_classified(gold_entries, classified_results):
    """Match gold standard entries to their predictions in the classified results.

    Uses (company_name, programme_name) tuple for matching to handle duplicate
    company names correctly (e.g. Howden has two programmes with different firm types).
    Normalizes text to handle encoding differences between xlsx and CSV sources.

    Returns:
        matched: list of (gold_entry, classified_entry) tuples
        unmatched_gold: list of gold entries with no match
        unseen: list of classified entries not in the gold standard
    """
    # Build lookup from classified results using normalized keys
    classified_lookup = {}
    for r in classified_results:
        key = (_normalize_for_matching(r["company_name"]),
               _normalize_for_matching(r["programme_name"]))
        classified_lookup[key] = r

    matched = []
    unmatched_gold = []

    matched_classified_keys = set()
    for g in gold_entries:
        key = (_normalize_for_matching(g["company_name"]),
               _normalize_for_matching(g["programme_name"]))

        if key in classified_lookup:
            matched.append((g, classified_lookup[key]))
            matched_classified_keys.add(key)
        else:
            unmatched_gold.append(g)

    # Unseen = classified entries not matched to any gold standard entry
    unseen = []
    for r in classified_results:
        key = (_normalize_for_matching(r["company_name"]),
               _normalize_for_matching(r["programme_name"]))
        if key not in matched_classified_keys:
            unseen.append(r)

    return matched, unmatched_gold, unseen


# ---------------------------------------------------------------------------
# 2. Per-class metrics (precision, recall, F1)
# ---------------------------------------------------------------------------

def evaluate_classifier(predictions, gold_standard, label=""):
    """Compute per-class precision, recall, F1 for a classification dimension.

    Args:
        predictions: list of predicted label strings.
        gold_standard: list of true label strings.
        label: descriptive name for printing.

    Returns:
        Dict with macro/weighted F1 and the full classification report,
        or None if gold_standard is None.
    """
    if gold_standard is None or len(gold_standard) == 0:
        print(f"[evaluate] No data for {label}, skipping.")
        return None

    all_labels = sorted(set(predictions + gold_standard))

    report_str = classification_report(
        gold_standard, predictions, labels=all_labels, zero_division=0
    )
    report_dict = classification_report(
        gold_standard, predictions, labels=all_labels, zero_division=0, output_dict=True
    )

    if label:
        print(f"\n--- {label} ---")
        print(report_str)

    return {
        "label": label,
        "macro_f1": round(report_dict["macro avg"]["f1-score"], 4),
        "weighted_f1": round(report_dict["weighted avg"]["f1-score"], 4),
        "report": report_str,
    }


# ---------------------------------------------------------------------------
# 3. Confusion matrix heatmap
# ---------------------------------------------------------------------------

def confusion_matrix_plot(predictions, gold_standard, label="firm_type", save_path=None):
    """Generate and save a confusion matrix heatmap. Returns the cm array."""
    if gold_standard is None or len(gold_standard) == 0:
        return None

    all_labels = sorted(set(predictions + gold_standard))
    cm = confusion_matrix(gold_standard, predictions, labels=all_labels)

    fig, ax = plt.subplots(figsize=(max(8, len(all_labels)), max(6, len(all_labels) * 0.8)))
    im = ax.imshow(cm, interpolation="nearest", cmap=plt.cm.Blues)
    ax.set_title(f"Confusion Matrix: {label}")
    fig.colorbar(im, ax=ax)

    ax.set_xticks(range(len(all_labels)))
    ax.set_yticks(range(len(all_labels)))
    ax.set_xticklabels(all_labels, rotation=45, ha="right", fontsize=8)
    ax.set_yticklabels(all_labels, fontsize=8)
    ax.set_xlabel("Predicted")
    ax.set_ylabel("True")

    for i in range(len(all_labels)):
        for j in range(len(all_labels)):
            ax.text(j, i, str(cm[i, j]), ha="center", va="center",
                    color="white" if cm[i, j] > cm.max() / 2 else "black")

    plt.tight_layout()

    if save_path is None:
        os.makedirs(FIGURES_DIR, exist_ok=True)
        save_path = os.path.join(FIGURES_DIR, f"confusion_matrix_{label}.png")

    fig.savefig(save_path, dpi=150)
    print(f"  [evaluate] Saved confusion matrix to {save_path}")
    plt.close(fig)

    return cm


# ---------------------------------------------------------------------------
# 4. Confidence calibration analysis
# ---------------------------------------------------------------------------

def calibration_analysis(predictions, gold_standard, confidences, label=""):
    """Group predictions by confidence band and compute actual accuracy.

    Returns dict with bin stats, or None if gold_standard is None.
    """
    if gold_standard is None or len(gold_standard) == 0:
        return None

    bins = []

    for i in range(len(CALIBRATION_BINS) - 1):
        lo, hi = CALIBRATION_BINS[i], CALIBRATION_BINS[i + 1]
        if i < len(CALIBRATION_BINS) - 2:
            mask = [(lo <= c < hi) for c in confidences]
        else:
            mask = [(lo <= c <= hi) for c in confidences]

        bin_true = [t for t, m in zip(gold_standard, mask) if m]
        bin_pred = [p for p, m in zip(predictions, mask) if m]
        bin_conf = [c for c, m in zip(confidences, mask) if m]

        n = len(bin_true)
        if n == 0:
            bins.append({"range": f"{lo:.1f}-{hi:.1f}", "n": 0, "accuracy": 0.0, "avg_conf": 0.0})
            continue

        accuracy = sum(1 for t, p in zip(bin_true, bin_pred) if t == p) / n
        avg_conf = sum(bin_conf) / n

        bins.append({
            "range": f"{lo:.1f}-{hi:.1f}",
            "n": n,
            "accuracy": round(accuracy, 4),
            "avg_conf": round(avg_conf, 4),
        })

    if label:
        print(f"\n--- Calibration: {label} ---")
        for b in bins:
            print(f"  {b['range']}: n={b['n']}, accuracy={b['accuracy']:.2f}, avg_conf={b['avg_conf']:.2f}")

    return {"bins": bins}


# ---------------------------------------------------------------------------
# 5. Expected Calibration Error (ECE)
# ---------------------------------------------------------------------------

def calculate_ece(predictions, gold_standard, confidences):
    """Compute ECE = sum over bins of (bin_weight * |accuracy - avg_confidence|).

    Lower is better; 0.0 means perfectly calibrated.
    """
    if gold_standard is None or len(gold_standard) == 0:
        return None

    cal = calibration_analysis(predictions, gold_standard, confidences)
    if cal is None:
        return None

    total = len(gold_standard)
    ece = 0.0
    for b in cal["bins"]:
        if b["n"] > 0:
            ece += (b["n"] / total) * abs(b["accuracy"] - b["avg_conf"])

    return round(ece, 4)


# ---------------------------------------------------------------------------
# 6. Reliability diagram
# ---------------------------------------------------------------------------

def reliability_diagram(predictions, gold_standard, confidences, label="", save_path=None):
    """Plot a reliability diagram showing calibration quality."""
    if gold_standard is None or len(gold_standard) == 0:
        return None

    cal = calibration_analysis(predictions, gold_standard, confidences)
    if cal is None:
        return None

    ece = calculate_ece(predictions, gold_standard, confidences)

    avg_confs = [b["avg_conf"] for b in cal["bins"] if b["n"] > 0]
    accuracies = [b["accuracy"] for b in cal["bins"] if b["n"] > 0]

    fig, ax = plt.subplots(figsize=(6, 6))
    ax.plot([0, 1], [0, 1], "k--", label="Perfect calibration")
    ax.bar(avg_confs, accuracies, width=0.15, alpha=0.6, label="Model", edgecolor="black")
    ax.set_xlabel("Mean predicted confidence")
    ax.set_ylabel("Fraction of correct predictions")
    ax.set_title(f"Reliability Diagram: {label}\nECE = {ece:.4f}")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.legend()
    plt.tight_layout()

    if save_path is None:
        os.makedirs(FIGURES_DIR, exist_ok=True)
        save_path = os.path.join(FIGURES_DIR, f"reliability_{label}.png")

    fig.savefig(save_path, dpi=150)
    print(f"  [evaluate] Saved reliability diagram to {save_path}")
    plt.close(fig)


# ---------------------------------------------------------------------------
# 7. Error analysis: top confusion pairs
# ---------------------------------------------------------------------------

def error_analysis(y_true, y_pred, dimension, n=5):
    """Extract top-N confused label pairs from predictions.

    Prints a ranked table of (true_label -> predicted_label, count) for
    off-diagonal confusion matrix entries. For firm_type, specifically
    flags Elite Boutique <-> Middle Market IB confusion.
    """
    if not y_true or not y_pred:
        return

    all_labels = sorted(set(y_true + y_pred))
    cm = confusion_matrix(y_true, y_pred, labels=all_labels)

    # Collect off-diagonal (true, pred, count) triples
    errors = []
    for i, true_label in enumerate(all_labels):
        for j, pred_label in enumerate(all_labels):
            if i != j and cm[i, j] > 0:
                errors.append((true_label, pred_label, cm[i, j]))

    errors.sort(key=lambda x: x[2], reverse=True)

    print(f"\n--- Error Analysis: {dimension} (top {n} confusions) ---")
    if not errors:
        print("  No misclassifications found.")
        return

    for true_label, pred_label, count in errors[:n]:
        print(f"  {true_label} -> {pred_label}: {count} times")

    # Specific hypothesis test for firm_type: EB <-> MM IB
    if dimension == "firm_type":
        eb_idx = all_labels.index("ELITE_BOUTIQUE") if "ELITE_BOUTIQUE" in all_labels else None
        mm_idx = all_labels.index("MIDDLE_MARKET_IB") if "MIDDLE_MARKET_IB" in all_labels else None
        if eb_idx is not None and mm_idx is not None:
            eb_to_mm = cm[eb_idx, mm_idx]
            mm_to_eb = cm[mm_idx, eb_idx]
            total_eb = sum(cm[eb_idx, :])
            total_mm = sum(cm[mm_idx, :])
            print(f"\n  Hypothesis check: Elite Boutique <-> Middle Market IB")
            print(f"    EB misclassified as MM: {eb_to_mm}/{total_eb} ({eb_to_mm/max(total_eb,1)*100:.0f}%)")
            print(f"    MM misclassified as EB: {mm_to_eb}/{total_mm} ({mm_to_eb/max(total_mm,1)*100:.0f}%)")


# ---------------------------------------------------------------------------
# 8. Class distribution report
# ---------------------------------------------------------------------------

def print_class_distribution(gold_entries):
    """Print gold standard class frequencies per dimension. Flag sparse classes."""
    print(f"\n{'=' * 60}")
    print("GOLD STANDARD CLASS DISTRIBUTION")
    print(f"{'=' * 60}")

    for dimension, valid_list in [("firm_type", FIRM_TYPES),
                                   ("role_function", ROLE_FUNCTIONS),
                                   ("programme_status", PROGRAMME_STATUSES)]:
        counts = Counter(g[dimension] for g in gold_entries)
        total = sum(counts.values())
        print(f"\n  {dimension} (n={total}):")
        for label in valid_list:
            c = counts.get(label, 0)
            pct = c / total * 100 if total > 0 else 0
            flag = " ** SPARSE - per-class metrics unreliable" if 0 < c < 5 else ""
            flag = " (absent)" if c == 0 else flag
            print(f"    {label:<25s} {c:>4d}  ({pct:5.1f}%){flag}")


# ---------------------------------------------------------------------------
# 9. Unseen split descriptive statistics
# ---------------------------------------------------------------------------

def unseen_split_report(unseen_entries):
    """Print descriptive stats for the entries NOT in the gold standard.

    No accuracy metrics (no ground truth), but shows confidence distributions,
    escalation rates, and LLM vs baseline agreement.
    """
    if not unseen_entries:
        print("\n[evaluate] No unseen entries to report.")
        return

    n = len(unseen_entries)
    print(f"\n{'=' * 60}")
    print(f"UNSEEN SPLIT DESCRIPTIVE STATISTICS ({n} entries)")
    print(f"{'=' * 60}")

    # Escalation rate
    escalated = sum(1 for r in unseen_entries if r.get("escalated") in (True, "True", "TRUE"))
    print(f"\n  Escalation rate: {escalated}/{n} ({escalated/n*100:.1f}%)")

    # Model distribution
    models = Counter(r.get("model_used", "unknown") for r in unseen_entries)
    print(f"  Model distribution: {dict(models)}")

    # Confidence distributions per dimension
    for dim in ["firm_type", "role_function", "programme_status"]:
        conf_key = f"{dim}_confidence"
        confs = [r.get(conf_key, 0.5) for r in unseen_entries]
        print(f"\n  {dim} confidence (unseen):")
        print(f"    mean={np.mean(confs):.3f}, median={np.median(confs):.3f}, "
              f"min={np.min(confs):.3f}, max={np.max(confs):.3f}")

        # Bucket distribution
        for i in range(len(CALIBRATION_BINS) - 1):
            lo, hi = CALIBRATION_BINS[i], CALIBRATION_BINS[i + 1]
            if i < len(CALIBRATION_BINS) - 2:
                count = sum(1 for c in confs if lo <= c < hi)
            else:
                count = sum(1 for c in confs if lo <= c <= hi)
            print(f"    {lo:.1f}-{hi:.1f}: {count} ({count/n*100:.1f}%)")

    # LLM vs baseline agreement
    print(f"\n  LLM vs rule-based baseline agreement:")
    for dim, baseline_key in [("firm_type", "baseline_firm_type"),
                               ("role_function", "baseline_role_function"),
                               ("programme_status", "baseline_programme_status")]:
        agree = sum(1 for r in unseen_entries if r.get(dim) == r.get(baseline_key))
        print(f"    {dim}: {agree}/{n} ({agree/n*100:.1f}%)")


# ---------------------------------------------------------------------------
# 10. Full evaluation runner
# ---------------------------------------------------------------------------

def run_full_evaluation(gold_path=GOLD_STANDARD_PATH,
                        classified_path=CLASSIFIED_PATH):
    """Run complete evaluation: random baseline vs rule-based vs LLM.

    Loads gold standard xlsx (153 hand-labelled entries), loads classified
    results xlsx (501 pipeline outputs), matches the 153 by (company_name,
    programme_name), runs all metrics, and reports on the unseen 348.
    """
    print(f"\n{'=' * 60}")
    print("EVALUATION PIPELINE")
    print(f"{'=' * 60}")

    # --- Load gold standard ---
    gold_entries = load_gold_standard(gold_path)
    if gold_entries is None:
        return
    print(f"[evaluate] Loaded {len(gold_entries)} gold standard entries")

    # --- Class distribution ---
    print_class_distribution(gold_entries)

    # --- Load classified results ---
    classified = load_classified_results(classified_path)
    if classified is None:
        return
    print(f"\n[evaluate] Loaded {len(classified)} classified results")

    # --- Match ---
    matched, unmatched, unseen = match_gold_to_classified(gold_entries, classified)
    print(f"[evaluate] Matched: {len(matched)}, Unmatched gold: {len(unmatched)}, "
          f"Unseen (no ground truth): {len(unseen)}")

    if unmatched:
        print("[evaluate] WARNING: unmatched gold standard entries:")
        for g in unmatched:
            print(f"  {g['company_name']} | {g['programme_name']}")

    if len(matched) < 5:
        print("[evaluate] Too few matched entries for evaluation.")
        return

    # --- Build parallel arrays for evaluation ---
    gold_labels = [m[0] for m in matched]
    llm_preds = [m[1] for m in matched]

    # Build listings for baseline generation (from gold standard entries)
    listings = []
    for g in gold_labels:
        listings.append({
            "company_name": g["company_name"],
            "programme_name": g["programme_name"],
            "opening_date": g.get("opening_date", ""),
            "closing_date": g.get("closing_date", ""),
            "latest_stage": "",
        })

    # Compute class weights from gold standard for proportional random baseline
    class_weights = {}
    for dimension in ["firm_type", "role_function", "programme_status"]:
        class_weights[dimension] = dict(Counter(g[dimension] for g in gold_labels))

    # Generate baselines
    random.seed(42)  # reproducible
    random_preds = random_baseline(listings, class_weights=class_weights)
    # Rule-based baseline: use predictions already in classified_listings.xlsx
    # (these were generated by the same pipeline that produced the LLM results)

    # --- Build methods for evaluation ---
    # For rule-based, extract from the matched classified entries
    rule_preds = []
    for _, c in matched:
        rule_preds.append({
            "firm_type": c.get("baseline_firm_type", "UNKNOWN"),
            "firm_type_confidence": c.get("baseline_firm_type_confidence", 0.0),
            "role_function": c.get("baseline_role_function", "UNKNOWN"),
            "role_function_confidence": 0.0,  # baseline doesn't have role conf in xlsx
            "programme_status": c.get("baseline_programme_status", "UNKNOWN"),
            "programme_status_confidence": 0.0,
        })

    methods = [
        ("Random", random_preds),
        ("Rule-based", rule_preds),
        ("LLM", [c for _, c in matched]),
    ]

    os.makedirs(FIGURES_DIR, exist_ok=True)
    summaries = []

    for method_name, preds_data in methods:
        print(f"\n{'=' * 60}")
        print(f"EVALUATION: {method_name}")
        print(f"{'=' * 60}")

        summary = {"method": method_name}

        for dimension in ["firm_type", "role_function", "programme_status"]:
            y_true = [g[dimension] for g in gold_labels]
            y_pred = [p[dimension] for p in preds_data]
            conf_key = f"{dimension}_confidence"
            confs = [p.get(conf_key, 0.5) for p in preds_data]

            # UNKNOWN is now included as a valid class for role_function
            # (previously 23 UNKNOWN gold labels were excluded here)

            metrics = evaluate_classifier(y_pred, y_true, label=f"{method_name} - {dimension}")
            if metrics:
                summary[f"{dimension}_macro_f1"] = metrics["macro_f1"]
                summary[f"{dimension}_weighted_f1"] = metrics["weighted_f1"]

            cm = confusion_matrix_plot(y_pred, y_true, label=f"{method_name}_{dimension}")

            # Error analysis (only for LLM and rule-based, not random)
            if method_name != "Random" and cm is not None:
                error_analysis(y_true, y_pred, dimension)

            cal = calibration_analysis(y_pred, y_true, confs, label=f"{method_name} - {dimension}")
            if cal:
                ece = calculate_ece(y_pred, y_true, confs)
                summary[f"{dimension}_ece"] = ece
                reliability_diagram(y_pred, y_true, confs, label=f"{method_name}_{dimension}")

        summaries.append(summary)

    # --- Comparison table ---
    print(f"\n{'=' * 80}")
    print("COMPARISON TABLE")
    print(f"{'=' * 80}")
    header = (f"{'Method':<15} | {'Firm F1':>8} | {'Role F1':>8} | {'Status F1':>10} | "
              f"{'Firm ECE':>9} | {'Role ECE':>9} | {'Status ECE':>11}")
    print(header)
    print("-" * len(header))
    for s in summaries:
        print(
            f"{s['method']:<15} | "
            f"{s.get('firm_type_macro_f1', 0):.4f}   | "
            f"{s.get('role_function_macro_f1', 0):.4f}   | "
            f"{s.get('programme_status_macro_f1', 0):.4f}     | "
            f"{s.get('firm_type_ece', 0):.4f}    | "
            f"{s.get('role_function_ece', 0):.4f}    | "
            f"{s.get('programme_status_ece', 0):.4f}"
        )
    print(f"{'=' * 80}")

    # --- Unseen split ---
    unseen_split_report(unseen)

    return summaries


# ---------------------------------------------------------------------------
# 11. CLI entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Run evaluation pipeline")
    parser.add_argument("--gold", default=GOLD_STANDARD_PATH,
                        help="Path to gold standard xlsx")
    parser.add_argument("--classified", default=CLASSIFIED_PATH,
                        help="Path to classified results xlsx")
    args = parser.parse_args()
    run_full_evaluation(args.gold, args.classified)
