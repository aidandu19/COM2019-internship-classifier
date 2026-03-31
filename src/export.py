# export.py - Excel and CSV export with conditional formatting for COM2019
#
# Exports classified internship data to a formatted Excel workbook with:
# - Firm type colour-coded cells
# - Confidence-based text styling (bold/normal/italic)
# - Status colour-coding (green/red/yellow/grey)
# - Auto-fitted column widths and frozen header row

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Colour palettes
# ---------------------------------------------------------------------------

FIRM_TYPE_COLOURS = {
    "BULGE_BRACKET":      "1F4E79",  # dark blue
    "ELITE_BOUTIQUE":     "2E7D32",  # dark green
    "MIDDLE_MARKET_IB":   "4A90D9",  # medium blue
    "BUY_SIDE_PE":        "6A1B9A",  # purple
    "ASSET_MANAGEMENT":   "00695C",  # teal
    "HEDGE_FUND":         "B71C1C",  # dark red
    "QUANT_PROP":         "E65100",  # deep orange
    "ACCOUNTING":         "37474F",  # dark blue-grey
    "CONSULTING":         "4E342E",  # brown
    "OTHER":              "616161",  # grey
    "UNKNOWN":            "9E9E9E",  # light grey
}

STATUS_COLOURS = {
    "OPEN":         "C8E6C9",  # light green
    "CLOSED":       "FFCDD2",  # light red
    "NOT_YET_OPEN": "FFF9C4",  # light yellow
    "UNKNOWN":      "E0E0E0",  # light grey
}

# Confidence-based text formatting thresholds:
# >= 0.8: high confidence (bold green) - classifier is very sure
# 0.5-0.8: medium confidence (normal dark grey) - reasonable certainty
# < 0.5: low confidence (italic red) - treat with caution
CONF_HIGH = Font(bold=True, color="1B5E20")
CONF_MED = Font(bold=False, color="333333")
CONF_LOW = Font(italic=True, color="B71C1C")


# ---------------------------------------------------------------------------
# Column definitions: (Excel header, DataFrame column name)
# ---------------------------------------------------------------------------

COLUMNS = [
    ("Company Name",            "company_name"),
    ("Programme",               "programme_name"),
    ("Opening Date",            "opening_date"),
    ("Closing Date",            "closing_date"),
    ("Latest Stage",            "latest_stage"),
    # Baseline results
    ("Baseline: Firm Type",     "baseline_firm_type"),
    ("Baseline: Firm Conf",     "baseline_firm_type_confidence"),
    ("Baseline: Role",          "baseline_role_function"),
    ("Baseline: Status",        "baseline_programme_status"),
    # LLM results
    ("LLM: Firm Type",          "llm_firm_type"),
    ("LLM: Firm Conf",          "llm_firm_type_confidence"),
    ("LLM: Firm Rationale",     "llm_firm_type_rationale"),
    ("LLM: Role",               "llm_role_function"),
    ("LLM: Role Conf",          "llm_role_function_confidence"),
    ("LLM: Role Rationale",     "llm_role_function_rationale"),
    ("LLM: Status",             "llm_programme_status"),
    ("LLM: Status Conf",        "llm_programme_status_confidence"),
    ("LLM: Status Rationale",   "llm_programme_status_rationale"),
    # Model escalation tracking
    ("LLM: Model Used",         "llm_model_used"),
    ("LLM: Escalated",          "llm_escalated"),
]


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------

def _confidence_font(conf_value):
    """Return font style based on confidence score."""
    try:
        conf = float(conf_value)
    except (ValueError, TypeError):
        return CONF_LOW
    if conf >= 0.8:
        return CONF_HIGH
    if conf >= 0.5:
        return CONF_MED
    return CONF_LOW


def _firm_type_fill(firm_type):
    """Return cell fill colour for a firm type label."""
    colour = FIRM_TYPE_COLOURS.get(str(firm_type).strip(), "FFFFFF")
    return PatternFill(start_color=colour, end_color=colour, fill_type="solid")


def _status_fill(status):
    """Return cell fill colour for a programme status label."""
    colour = STATUS_COLOURS.get(str(status).strip(), "FFFFFF")
    return PatternFill(start_color=colour, end_color=colour, fill_type="solid")


# ---------------------------------------------------------------------------
# Main Excel export
# ---------------------------------------------------------------------------

def export_to_excel(df, output_dir="output", filename=None):
    """Export a classified DataFrame to a formatted Excel workbook.

    Applies conditional formatting for firm types, confidence scores,
    and programme statuses. Auto-fits column widths and freezes the header.

    Args:
        df: pandas DataFrame with baseline_* and optionally llm_* columns.
        output_dir: directory to save the file in.
        filename: optional filename. Defaults to classified_internships.xlsx.

    Returns:
        Path to the saved file.
    """
    os.makedirs(output_dir, exist_ok=True)

    if filename is None:
        filename = "classified_internships.xlsx"

    filepath = os.path.join(output_dir, filename)

    # Filter columns to only those present in the DataFrame
    active_columns = [(header, key) for header, key in COLUMNS if key in df.columns]

    wb = Workbook()
    ws = wb.active
    ws.title = "Classifications"

    # --- Header row ---
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_idx, (header_name, _) in enumerate(active_columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # --- Data rows ---
    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for col_idx, (_, key) in enumerate(active_columns, start=1):
            value = row.get(key, "")
            if value is None:
                value = ""

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border

            # Firm type colour coding (white text on coloured background)
            if key in ("llm_firm_type", "baseline_firm_type"):
                cell.fill = _firm_type_fill(value)
                cell.font = Font(color="FFFFFF", bold=True)

            # Status colour coding
            if key in ("llm_programme_status", "baseline_programme_status"):
                cell.fill = _status_fill(value)

            # Confidence-based font styling
            if "confidence" in key:
                cell.font = _confidence_font(value)
                cell.alignment = Alignment(horizontal="center")

    # --- Auto-fit column widths ---
    for col_idx in range(1, len(active_columns) + 1):
        max_width = len(str(ws.cell(row=1, column=col_idx).value or ""))
        for row_idx in range(2, min(ws.max_row + 1, 52)):
            cell_len = len(str(ws.cell(row=row_idx, column=col_idx).value or ""))
            if cell_len > max_width:
                max_width = cell_len
        # Cap at 45 chars to prevent absurdly wide rationale columns
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 3, 45)

    # --- Freeze header row ---
    ws.freeze_panes = "A2"

    wb.save(filepath)
    print(f"[export] Saved Excel workbook to {filepath}")
    return filepath


# ---------------------------------------------------------------------------
# CSV fallback export
# ---------------------------------------------------------------------------

def export_to_csv(df, output_dir="output", filename=None):
    """Export a classified DataFrame to CSV (fallback if openpyxl fails)."""
    os.makedirs(output_dir, exist_ok=True)

    if filename is None:
        filename = "classified_internships.csv"

    filepath = os.path.join(output_dir, filename)
    df.to_csv(filepath, index=False)
    print(f"[export] Saved CSV to {filepath}")
    return filepath
