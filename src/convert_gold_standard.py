# convert_gold_standard.py - Convert Excel short codes to full enum names
#
# Reads the hand-labelled gold standard Excel file, converts short codes
# (BB, C, IB, etc.) to full taxonomy enum names (BULGE_BRACKET, CLOSED, etc.),
# writes data/gold_standard.json and updates the Excel in-place.

import json
import sys
from collections import Counter

from openpyxl import load_workbook

from src.taxonomy import (
    FIRM_TYPES, ROLE_FUNCTIONS, PROGRAMME_STATUSES,
    FIRM_TYPE_CODES, ROLE_FUNCTION_CODES, PROGRAMME_STATUS_CODES,
)

EXCEL_PATH = "data/GOLD STANDARD HAND FILLED LISTINGS.xlsx"
JSON_PATH = "data/gold_standard.json"
SHEET_NAME = "Gold Standard"

# Column positions (1-indexed) in the Excel sheet
COL_COMPANY = 2
COL_PROGRAMME = 3
COL_OPENING = 4
COL_CLOSING = 5
COL_STAGE = 6
COL_C_FT = 10   # AI-generated firm type
COL_C_ST = 11   # AI-generated programme status
COL_C_RF = 12   # AI-generated role function
COL_YOUR_FT = 15  # Hand-labelled firm type
COL_YOUR_ST = 16  # Hand-labelled programme status
COL_YOUR_RF = 17  # Hand-labelled role function
COL_RATIONALE = 18


def convert_code(code, mapping, valid_list, dimension_name, row_id, company):
    """Look up a short code in the mapping dict and validate against taxonomy.

    Returns the full enum name, or None with a printed warning if conversion fails.
    """
    if code is None or str(code).strip() == "":
        return None

    code_str = str(code).strip().upper()

    # Already a full enum name?
    if code_str in valid_list:
        return code_str

    # Look up in mapping
    full_name = mapping.get(code_str)
    if full_name is None:
        print(f"  WARNING row {row_id} ({company}): unknown {dimension_name} code '{code_str}'")
        return None

    if full_name not in valid_list:
        print(f"  WARNING row {row_id} ({company}): mapped {dimension_name} '{full_name}' not in taxonomy")
        return None

    return full_name


def main():
    print(f"Loading {EXCEL_PATH}...")
    wb = load_workbook(EXCEL_PATH)

    if SHEET_NAME not in wb.sheetnames:
        print(f"ERROR: Sheet '{SHEET_NAME}' not found. Available: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[SHEET_NAME]

    entries = []
    skipped = []
    ft_counter = Counter()
    st_counter = Counter()
    rf_counter = Counter()

    for row_num in range(2, 155):  # rows 2 through 154
        company = ws.cell(row=row_num, column=COL_COMPANY).value
        if company is None or str(company).strip() == "":
            continue

        company = str(company).strip()
        programme = str(ws.cell(row=row_num, column=COL_PROGRAMME).value or "").strip()
        opening_date = str(ws.cell(row=row_num, column=COL_OPENING).value or "").strip()
        closing_date = str(ws.cell(row=row_num, column=COL_CLOSING).value or "").strip()
        latest_stage = str(ws.cell(row=row_num, column=COL_STAGE).value or "").strip()

        # Convert YOUR_ columns (hand labels)
        your_ft_raw = ws.cell(row=row_num, column=COL_YOUR_FT).value
        your_st_raw = ws.cell(row=row_num, column=COL_YOUR_ST).value
        your_rf_raw = ws.cell(row=row_num, column=COL_YOUR_RF).value

        your_ft = convert_code(your_ft_raw, FIRM_TYPE_CODES, FIRM_TYPES,
                               "firm_type", row_num, company)
        your_st = convert_code(your_st_raw, PROGRAMME_STATUS_CODES, PROGRAMME_STATUSES,
                               "programme_status", row_num, company)
        your_rf = convert_code(your_rf_raw, ROLE_FUNCTION_CODES, ROLE_FUNCTIONS,
                               "role_function", row_num, company)

        if your_ft is None or your_st is None or your_rf is None:
            skipped.append((row_num, company))
            continue

        # Convert c_ columns (AI labels) - best effort, don't skip if these fail
        c_ft_raw = ws.cell(row=row_num, column=COL_C_FT).value
        c_st_raw = ws.cell(row=row_num, column=COL_C_ST).value
        c_rf_raw = ws.cell(row=row_num, column=COL_C_RF).value

        c_ft = convert_code(c_ft_raw, FIRM_TYPE_CODES, FIRM_TYPES,
                            "c_firm_type", row_num, company)
        c_st = convert_code(c_st_raw, PROGRAMME_STATUS_CODES, PROGRAMME_STATUSES,
                            "c_programme_status", row_num, company)
        c_rf = convert_code(c_rf_raw, ROLE_FUNCTION_CODES, ROLE_FUNCTIONS,
                            "c_role_function", row_num, company)

        # Build entry
        entry = {
            "company_name": company,
            "programme": programme,
            "opening_date": opening_date,
            "closing_date": closing_date,
            "latest_stage": latest_stage,
            "firm_type": your_ft,
            "programme_status": your_st,
            "role_function": your_rf,
        }
        entries.append(entry)

        # Track distribution
        ft_counter[your_ft] += 1
        st_counter[your_st] += 1
        rf_counter[your_rf] += 1

        # Overwrite Excel cells with full enum names
        ws.cell(row=row_num, column=COL_YOUR_FT, value=your_ft)
        ws.cell(row=row_num, column=COL_YOUR_ST, value=your_st)
        ws.cell(row=row_num, column=COL_YOUR_RF, value=your_rf)

        if c_ft:
            ws.cell(row=row_num, column=COL_C_FT, value=c_ft)
        if c_st:
            ws.cell(row=row_num, column=COL_C_ST, value=c_st)
        if c_rf:
            ws.cell(row=row_num, column=COL_C_RF, value=c_rf)

    # Write JSON
    with open(JSON_PATH, "w") as f:
        json.dump(entries, f, indent=2)
    print(f"\nWrote {len(entries)} entries to {JSON_PATH}")

    # Save Excel
    wb.save(EXCEL_PATH)
    print(f"Updated Excel in-place: {EXCEL_PATH}")

    # Summary
    print(f"\n{'=' * 50}")
    print(f"CONVERSION SUMMARY")
    print(f"{'=' * 50}")
    print(f"Total converted: {len(entries)}")
    print(f"Skipped:         {len(skipped)}")

    if skipped:
        print(f"\nSkipped rows:")
        for row_num, company in skipped:
            print(f"  Row {row_num}: {company}")

    print(f"\nFirm Type distribution:")
    for ft, count in ft_counter.most_common():
        print(f"  {ft}: {count}")

    print(f"\nProgramme Status distribution:")
    for st, count in st_counter.most_common():
        print(f"  {st}: {count}")

    print(f"\nRole Function distribution:")
    for rf, count in rf_counter.most_common():
        print(f"  {rf}: {count}")


if __name__ == "__main__":
    main()
