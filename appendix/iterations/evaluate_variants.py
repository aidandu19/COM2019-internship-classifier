# evaluate_variants.py - Compare V1 vs V2 prompts and Schema A vs Schema B
#
# Loads gold standard labels and LLM results from multiple experiment files,
# computes F1 and ECE for each, and produces comparison tables and confusion
# matrices. Used for COM2019 Tasks 6c and 6d.

import os
import json
import sys
from collections import Counter

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from sklearn.metrics import classification_report, confusion_matrix

from src.taxonomy import FIRM_TYPES, ROLE_FUNCTIONS, PROGRAMME_STATUSES
from src.taxonomy import FIRM_TYPE_CODES, ROLE_FUNCTION_CODES, PROGRAMME_STATUS_CODES

FIGURES_DIR = os.path.join("output", "figures")
CALIBRATION_BINS = [0.0, 0.3, 0.6, 0.8, 1.0]


def load_gold_labels():
    """Load gold standard labels from the Excel file (same as evaluate.py)."""
    from openpyxl import load_workbook
    path = "data/GOLD STANDARD HAND FILLED LISTINGS.xlsx"
    wb = load_workbook(path)
    ws = wb["Gold Standard"]
    entries = []
    for row_idx in range(2, ws.max_row + 1):
        company = ws.cell(row=row_idx, column=2).value
        if company is None or str(company).strip() == "":
            continue
        company = str(company).strip()
        programme = str(ws.cell(row=row_idx, column=3).value or "").strip()
        ft_raw = ws.cell(row=row_idx, column=7).value
        st_raw = ws.cell(row=row_idx, column=6).value
        rf_raw = ws.cell(row=row_idx, column=8).value

        def convert(code, mapping, valid):
            if code is None or str(code).strip() == "":
                return None
            s = str(code).strip().upper()
            if s in valid:
                return s
            return mapping.get(s)

        ft = convert(ft_raw, FIRM_TYPE_CODES, FIRM_TYPES)
        st = convert(st_raw, PROGRAMME_STATUS_CODES, PROGRAMME_STATUSES)
        rf = convert(rf_raw, ROLE_FUNCTION_CODES, ROLE_FUNCTIONS)
        if ft is None or st is None or rf is None:
            continue
        entries.append({
            "company_name": company,
            "programme_name": programme,
            "firm_type": ft,
            "programme_status": st,
            "role_function": rf,
        })
    wb.close()
    return entries


def load_llm_results(path):
    """Load LLM results from a JSON file."""
    with open(path) as f:
        data = json.load(f)
    return data


def load_gold_standard_json():
    """Load gold_standard.json to get company ordering for matching."""
    with open("data/gold_standard.json") as f:
        data = json.load(f)
    return data


def _norm(text):
    import unicodedata
    text = unicodedata.normalize("NFKD", str(text))
    text = "".join(c for c in text if not unicodedata.combining(c))
    for ch in ("\u2013", "\u2014", "\u00d0"):
        text = text.replace(ch, "-")
    return text.strip().lower()


def match_by_name(gold_labels, llm_results):
    """Match LLM results to gold labels by (company_name, programme_name).

    Works whether llm_results has 153 entries (gold-only run) or 501 (full TRACKR).
    """
    # Build lookup from LLM results
    llm_lookup = {}
    for lr in llm_results:
        company = lr.get("company_name", "")
        programme = lr.get("programme_name") or lr.get("programme", "")
        key = (_norm(company), _norm(programme))
        llm_lookup[key] = lr

    matched_gold = []
    matched_llm = []
    for g in gold_labels:
        key = (_norm(g["company_name"]), _norm(g["programme_name"]))
        if key in llm_lookup:
            matched_gold.append(g)
            matched_llm.append(llm_lookup[key])

    return matched_gold, matched_llm


def match_variant_results(gold_labels, llm_results, gold_json):
    """Match variant results (V2/Schema B) that are in gold_standard.json order."""
    # These results are in the same order as gold_standard.json (153 entries)
    matched_gold = []
    matched_llm = []

    # Build lookup from gold labels
    gold_lookup = {}
    for g in gold_labels:
        key = (_norm(g["company_name"]), _norm(g["programme_name"]))
        gold_lookup[key] = g

    for gj, lr in zip(gold_json, llm_results):
        gj_key = (_norm(gj.get("company_name", "")),
                   _norm(gj.get("programme") or gj.get("programme_name", "")))
        if gj_key in gold_lookup:
            matched_gold.append(gold_lookup[gj_key])
            matched_llm.append(lr)

    return matched_gold, matched_llm


def calculate_ece(y_true, y_pred, confs):
    """Compute Expected Calibration Error."""
    total = len(y_true)
    if total == 0:
        return 0.0
    ece = 0.0
    for i in range(len(CALIBRATION_BINS) - 1):
        lo, hi = CALIBRATION_BINS[i], CALIBRATION_BINS[i + 1]
        if i < len(CALIBRATION_BINS) - 2:
            mask = [(lo <= c < hi) for c in confs]
        else:
            mask = [(lo <= c <= hi) for c in confs]
        bin_true = [t for t, m in zip(y_true, mask) if m]
        bin_pred = [p for p, m in zip(y_pred, mask) if m]
        bin_conf = [c for c, m in zip(confs, mask) if m]
        n = len(bin_true)
        if n == 0:
            continue
        accuracy = sum(1 for t, p in zip(bin_true, bin_pred) if t == p) / n
        avg_conf = sum(bin_conf) / n
        ece += (n / total) * abs(accuracy - avg_conf)
    return round(ece, 4)


def evaluate_dimension(gold_entries, llm_entries, dimension, label_prefix=""):
    """Evaluate one dimension, return metrics dict."""
    y_true = [g[dimension] for g in gold_entries]
    y_pred = [p[dimension] for p in llm_entries]
    conf_key = f"{dimension}_confidence"
    confs = [float(p.get(conf_key, 0.5)) for p in llm_entries]

    all_labels = sorted(set(y_true + y_pred))
    report = classification_report(y_true, y_pred, labels=all_labels,
                                   zero_division=0, output_dict=True)

    macro_f1 = round(report["macro avg"]["f1-score"], 4)
    ece = calculate_ece(y_true, y_pred, confs)

    return {
        "macro_f1": macro_f1,
        "ece": ece,
        "y_true": y_true,
        "y_pred": y_pred,
        "confs": confs,
        "report_str": classification_report(y_true, y_pred, labels=all_labels, zero_division=0),
    }


def plot_confusion_matrix(y_true, y_pred, label, save_path):
    """Generate confusion matrix plot."""
    all_labels = sorted(set(y_true + y_pred))
    cm = confusion_matrix(y_true, y_pred, labels=all_labels)
    fig, ax = plt.subplots(figsize=(max(8, len(all_labels)), max(6, len(all_labels) * 0.8)))
    im = ax.imshow(cm, interpolation="nearest", cmap=plt.cm.Blues)
    ax.set_title(f"Confusion Matrix: {label}")
    fig.colorbar(im, ax=ax)
    ax.set_xticks(range(len(all_labels)))
    ax.set_yticks(range(len(all_labels)))
    ax.set_xticklabels(all_labels, rotation=45, ha="right", fontsize=8)
    ax.set_yticklabels(all_labels, fontsize=8)
    ax.set_xlabel("Predicted")
    ax.set_ylabel("True")
    for i in range(len(all_labels)):
        for j in range(len(all_labels)):
            ax.text(j, i, str(cm[i, j]), ha="center", va="center",
                    color="white" if cm[i, j] > cm.max() / 2 else "black")
    plt.tight_layout()
    os.makedirs(os.path.dirname(save_path), exist_ok=True)
    fig.savefig(save_path, dpi=150)
    plt.close(fig)
    print(f"  Saved: {save_path}")


def load_v1_from_classified_xlsx():
    """Load V1 LLM predictions from classified_listings.xlsx (same as evaluate.py)."""
    from openpyxl import load_workbook
    path = "data/classified_listings.xlsx"
    wb = load_workbook(path)
    ws = wb.active
    results = []
    for row_idx in range(2, ws.max_row + 1):
        company = ws.cell(row=row_idx, column=1).value
        if company is None or str(company).strip() == "":
            continue
        def _cell(col):
            v = ws.cell(row=row_idx, column=col).value
            return v if v is not None else ""
        def _float(col):
            v = ws.cell(row=row_idx, column=col).value
            try:
                return float(v)
            except (ValueError, TypeError):
                return 0.5
        results.append({
            "company_name": str(company).strip(),
            "programme_name": str(_cell(2)).strip(),
            "firm_type": str(_cell(10)).strip(),
            "firm_type_confidence": _float(11),
            "firm_type_rationale": str(_cell(12)).strip(),
            "role_function": str(_cell(13)).strip(),
            "role_function_confidence": _float(14),
            "role_function_rationale": str(_cell(15)).strip(),
            "programme_status": str(_cell(16)).strip(),
            "programme_status_confidence": _float(17),
            "programme_status_rationale": str(_cell(18)).strip(),
        })
    wb.close()
    return results


def run_comparison(experiment_name, results_path):
    """Run full comparison between V1 and an experiment variant."""
    print(f"\n{'=' * 70}")
    print(f"EVALUATION: {experiment_name}")
    print(f"{'=' * 70}")

    gold_labels = load_gold_labels()
    gold_json = load_gold_standard_json()
    print(f"Gold standard: {len(gold_labels)} entries")

    # Load V1 from classified_listings.xlsx (has programme_name for matching)
    v1_results = load_v1_from_classified_xlsx()
    print(f"V1 classified results: {len(v1_results)} entries")

    # Load variant results
    variant_results = load_llm_results(results_path)
    print(f"{experiment_name} results: {len(variant_results)} entries")

    # Match V1 by name (501 entries -> 153 matched), variant by position (153 entries)
    v1_gold, v1_llm = match_by_name(gold_labels, v1_results)
    var_gold, var_llm = match_variant_results(gold_labels, variant_results, gold_json)
    print(f"V1 matched: {len(v1_gold)}, {experiment_name} matched: {len(var_gold)}")

    # Evaluate all dimensions for both
    dimensions = ["firm_type", "role_function", "programme_status"]
    v1_metrics = {}
    var_metrics = {}

    for dim in dimensions:
        v1_metrics[dim] = evaluate_dimension(v1_gold, v1_llm, dim, "V1")
        var_metrics[dim] = evaluate_dimension(var_gold, var_llm, dim, experiment_name)

    # Print comparison table
    print(f"\n{'=' * 80}")
    print(f"COMPARISON: V1 (Schema A) vs {experiment_name}")
    print(f"{'=' * 80}")
    header = f"{'Metric':<20} | {'V1 (baseline)':<15} | {experiment_name:<15} | {'Delta':<10}"
    print(header)
    print("-" * len(header))

    for dim in dimensions:
        v1_f1 = v1_metrics[dim]["macro_f1"]
        var_f1 = var_metrics[dim]["macro_f1"]
        delta_f1 = var_f1 - v1_f1
        sign = "+" if delta_f1 >= 0 else ""
        print(f"{dim + ' F1':<20} | {v1_f1:<15.4f} | {var_f1:<15.4f} | {sign}{delta_f1:.4f}")

    for dim in dimensions:
        v1_ece = v1_metrics[dim]["ece"]
        var_ece = var_metrics[dim]["ece"]
        delta_ece = var_ece - v1_ece
        sign = "+" if delta_ece >= 0 else ""
        print(f"{dim + ' ECE':<20} | {v1_ece:<15.4f} | {var_ece:<15.4f} | {sign}{delta_ece:.4f}")

    print(f"{'=' * 80}")

    # Print detailed classification reports for the variant
    for dim in dimensions:
        print(f"\n--- {experiment_name} {dim} ---")
        print(var_metrics[dim]["report_str"])

    # Generate confusion matrices for the variant
    os.makedirs(FIGURES_DIR, exist_ok=True)
    for dim in dimensions:
        save_path = os.path.join(FIGURES_DIR,
                                 f"confusion_matrix_{experiment_name}_{dim}.png")
        plot_confusion_matrix(
            var_metrics[dim]["y_true"],
            var_metrics[dim]["y_pred"],
            f"{experiment_name}_{dim}",
            save_path,
        )

    # Error analysis for programme_status (key dimension for V2)
    print(f"\n--- Error Analysis: {experiment_name} programme_status ---")
    ps_true = var_metrics["programme_status"]["y_true"]
    ps_pred = var_metrics["programme_status"]["y_pred"]
    all_labels = sorted(set(ps_true + ps_pred))
    cm = confusion_matrix(ps_true, ps_pred, labels=all_labels)
    errors = []
    for i, tl in enumerate(all_labels):
        for j, pl in enumerate(all_labels):
            if i != j and cm[i, j] > 0:
                errors.append((tl, pl, cm[i, j]))
    errors.sort(key=lambda x: x[2], reverse=True)
    for tl, pl, count in errors[:10]:
        print(f"  {tl} -> {pl}: {count} times")
    if not errors:
        print("  No misclassifications!")

    return v1_metrics, var_metrics


def compare_rationale_examples(variant_path, gold_json, gold_labels, n=5):
    """Compare rationale quality between V1 (from classified xlsx) and variant."""
    v1_results = load_v1_from_classified_xlsx()
    var_results = load_llm_results(variant_path)

    # Build V1 lookup by name
    v1_lookup = {}
    for v1r in v1_results:
        key = (_norm(v1r.get("company_name", "")),
               _norm(v1r.get("programme_name", "")))
        v1_lookup[key] = v1r

    gold_lookup = {}
    for g in gold_labels:
        key = (_norm(g["company_name"]), _norm(g["programme_name"]))
        gold_lookup[key] = g

    print(f"\n{'=' * 70}")
    print(f"RATIONALE COMPARISON (interesting examples)")
    print(f"{'=' * 70}")

    examples_found = 0
    for gj, varr in zip(gold_json, var_results):
        gj_key = (_norm(gj.get("company_name", "")),
                   _norm(gj.get("programme") or gj.get("programme_name", "")))
        if gj_key not in gold_lookup or gj_key not in v1_lookup:
            continue
        gold = gold_lookup[gj_key]
        v1r = v1_lookup[gj_key]

        for dim in ["firm_type", "role_function", "programme_status"]:
            v1_pred = v1r.get(dim)
            var_pred = varr.get(dim)
            true_label = gold[dim]

            if v1_pred != true_label and (var_pred != v1_pred):
                examples_found += 1
                v1_conf = v1r.get(f"{dim}_confidence", "?")
                var_conf = varr.get(f"{dim}_confidence", "?")
                print(f"\n  Example {examples_found}: {gj.get('company_name')} ({dim})")
                print(f"    Gold:    {true_label}")
                print(f"    V1 pred: {v1_pred} (conf={v1_conf})")
                print(f"    Var pred:{var_pred} (conf={var_conf})")
                print(f"    V1 rationale:  {v1r.get(dim + '_rationale', 'N/A')[:120]}")
                print(f"    Var rationale: {varr.get(dim + '_rationale', 'N/A')[:120]}")

                changed = "FIXED" if var_pred == true_label else "STILL WRONG"
                print(f"    Result: {changed}")

                if examples_found >= n:
                    return


def main():
    if len(sys.argv) < 2:
        print("Usage:")
        print("  python evaluate_variants.py v2       # Compare V1 vs V2 prompt")
        print("  python evaluate_variants.py schema_b  # Compare Schema A vs Schema B")
        print("  python evaluate_variants.py all       # Run all comparisons")
        sys.exit(1)

    mode = sys.argv[1].lower()

    if mode in ("v2", "all"):
        v2_path = "data/llm_results_v2.json"
        if os.path.exists(v2_path):
            v1m, v2m = run_comparison("Prompt_V2", v2_path)

            # Print the specific V1 vs V2 vs Rule-based table for programme_status
            print(f"\n{'=' * 70}")
            print("PROGRAMME STATUS: V1 vs V2 vs RULE-BASED")
            print(f"{'=' * 70}")
            print(f"| {'Metric':<15} | {'Prompt V1':<12} | {'Prompt V2':<12} | {'Rule-based':<12} |")
            print(f"|{'-'*17}|{'-'*14}|{'-'*14}|{'-'*14}|")
            print(f"| {'Status F1':<15} | {v1m['programme_status']['macro_f1']:<12.4f} | {v2m['programme_status']['macro_f1']:<12.4f} | {'0.9949':<12} |")
            print(f"| {'Status ECE':<15} | {v1m['programme_status']['ece']:<12.4f} | {v2m['programme_status']['ece']:<12.4f} | {'N/A':<12} |")
        else:
            print(f"V2 results not found at {v2_path}. Run classify_v2.py first.")

    if mode in ("schema_b", "all"):
        sb_path = "data/llm_results_schema_b.json"
        if os.path.exists(sb_path):
            v1m, sbm = run_comparison("Schema_B", sb_path)

            # Rationale comparison
            gold_labels = load_gold_labels()
            gold_json = load_gold_standard_json()
            compare_rationale_examples(sb_path, gold_json, gold_labels, n=5)
        else:
            print(f"Schema B results not found at {sb_path}. Run classify_schema_b.py first.")


if __name__ == "__main__":
    main()
